"""为测试报告的每张图保存中文说明和完整关键数值表。"""

from __future__ import annotations

from importlib import metadata as importlib_metadata
from pathlib import Path
import sys
from typing import Any


def _distribution_version(name: str) -> str | None:
    try:
        return importlib_metadata.version(name)
    except importlib_metadata.PackageNotFoundError:
        return None


def _word_dependency_message(detail: str) -> str:
    executable = f'& "{sys.executable}"' if sys.platform == "win32" else f'"{sys.executable}"'
    return (
        f"Word 文档依赖不可用：{detail}\n\n"
        "请使用运行本脚本的同一个 Python 解释器执行：\n"
        f"  {executable} -m pip uninstall -y docx\n"
        f"  {executable} -m pip install --upgrade --force-reinstall "
        '"python-docx>=1.1,<2" "lxml>=4.9,<7"\n\n'
        "不要安装 exceptions，也不要安装名为 docx 的旧包。"
    )


_legacy_docx = _distribution_version("docx")
_python_docx = _distribution_version("python-docx")
if _legacy_docx is not None:
    raise SystemExit(_word_dependency_message(
        f"检测到旧包 docx=={_legacy_docx}；它会在 Python 3 中触发 No module named 'exceptions'。"))
if _python_docx is None:
    raise SystemExit(_word_dependency_message("未安装 python-docx。"))
try:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
except Exception as error:
    hint = "当前 docx 模块仍在导入 Python 2 的 exceptions。" if (
        isinstance(error, ModuleNotFoundError) and error.name == "exceptions"
    ) else f"导入 python-docx 失败：{type(error).__name__}: {error}"
    raise SystemExit(_word_dependency_message(hint)) from error


def _font(style: Any, name: str, size: float) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)


def _text(value: Any) -> str:
    if value is None:
        return "缺失"
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, float):
        return format(value, ".10g")
    return str(value)


def _add_table(document: Any, headers: list[str], rows: list[list[Any]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement("w:" + edge)
        for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            element.set(qn("w:" + key), value)
        borders.append(element)
    table._tbl.tblPr.append(borders)
    # 按内容长度分配列宽；限制长路径等离群内容对列宽的影响。
    lengths = [
        min(24, max(6, len(str(header)), *(len(_text(row[i])) for row in rows)))
        for i, header in enumerate(headers)
    ]
    widths = [Cm(17 * length / sum(lengths)) for length in lengths]
    for column, width in zip(table.columns, widths):
        column.width = width
    for cell, width, header in zip(table.rows[0].cells, widths, headers):
        cell.width = width
        cell.text = str(header)
        shade = OxmlElement("w:shd")
        shade.set(qn("w:fill"), "DCEAF3")
        cell._tc.get_or_add_tcPr().append(shade)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    repeat = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for values in rows:
        for cell, width, value in zip(table.add_row().cells, widths, values):
            cell.width = width
            cell.text = _text(value)
    for row in table.rows:
        # 表格可自然续页，单行不得被截成两页，且不使用固定行高。
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.paragraph_format.space_before = Pt(3)
                paragraph.paragraph_format.keep_with_next = False
                for run in paragraph.runs:
                    run.font.size = Pt(10)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def write_chart_document(
    path: str | Path,
    title: str,
    section: str,
    caption: str,
    notes: list[str],
    tables: list[dict[str, Any]],
    sources: list[str],
) -> None:
    """写入一张图的 DOCX，完整保留传入的表格行和数值。

    ``tables`` 每项含 ``title``、``headers``、``rows``。调用方负责选择
    关键数据；本函数不抽样、不省略尾部，小数显示十位有效数字。
    """
    for item in tables:
        headers = item["headers"]
        if not headers or any(len(row) != len(headers) for row in item["rows"]):
            raise ValueError(f"数据表列数不一致或缺少表头：{item['title']}")

    document = Document()
    page = document.sections[0]
    page.page_width, page.page_height = Cm(21), Cm(29.7)
    page.top_margin = page.bottom_margin = Cm(2)
    page.left_margin = page.right_margin = Cm(2)
    for name, font, size in (
        ("Normal", "宋体", 11),
        ("Title", "黑体", 18),
        ("Heading 1", "黑体", 13),
        ("Heading 2", "黑体", 11),
    ):
        _font(document.styles[name], font, size)
    document.styles["Normal"].paragraph_format.space_after = Pt(6)
    document.styles["Normal"].paragraph_format.line_spacing = 1.15
    document.core_properties.title = title
    document.core_properties.subject = "图表关键数据与统计说明"
    document.add_paragraph(f"{title} 关键数据", style="Title")
    document.add_paragraph(f"对应测试报告章节：{section}")
    document.add_paragraph("本文件列出该图的图注、统计说明和关键数据，用于核对图中数值及复现图表。")
    document.add_paragraph("小数显示十位有效数字，未舍入的完整数值保存在同目录的数据表和数据文件中。")
    document.add_heading("图注", level=1)
    document.add_paragraph(caption)
    if notes:
        document.add_heading("统计说明", level=1)
        for note in notes:
            document.add_paragraph(str(note))
    for item in tables:
        document.add_heading(str(item["title"]), level=1)
        headers, rows = item["headers"], item["rows"]
        if not rows:
            document.add_paragraph("暂无可用数据。")
            continue
        # 宽表按列分组，并保留首列作为行标识，所有数据均被保留。
        if len(headers) <= 6:
            _add_table(document, headers, rows)
        else:
            for start in range(1, len(headers), 5):
                indices = [0, *range(start, min(start + 5, len(headers)))]
                document.add_heading(f"第 {(start - 1) // 5 + 1} 组数据列", level=2)
                _add_table(document, [headers[i] for i in indices], [[row[i] for i in indices] for row in rows])
    if sources:
        document.add_heading("数据来源", level=1)
        for source in dict.fromkeys(sources):
            paragraph = document.add_paragraph(str(source))
            paragraph.paragraph_format.space_after = Pt(3)
            for run in paragraph.runs:
                run.font.size = Pt(8)
    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output))
